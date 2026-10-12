"""Solution-EDV.xlsx export/import.

Export patches cell XML inside the original template (assets/Solution-EDV.xlsx) so every style,
merged range, dropdown list and Excel cell-checkbox survives untouched (openpyxl would drop the
checkbox feature bag). Import reads values with openpyxl.

Row keys mirror frontend/src/lib/solutionRows.ts — keep both in sync.
"""
import io
import re
import zipfile
from pathlib import Path
from typing import Dict, Optional
from xml.sax.saxutils import escape

import openpyxl

from models.client import SolutionCell, SolutionSystem, SolutionTable

TEMPLATE = Path(__file__).resolve().parent.parent / "assets" / "Solution-EDV.xlsx"
SHEET_PATH = "xl/worksheets/sheet1.xml"

VALUE_COLS = ["C", "E", "G", "I", "K", "M"]  # one per profile system
CHECK_COLS = ["D", "F", "H", "J", "L", "N"]  # checkbox / second value column
NAME_ROW = 1
COLOR_ROWS = range(3, 8)  # Culoare capace: 5 rows x 2 cols = up to 10 colours per system

# key -> (row, mode) ; mode: "value" | "value_check" | "check" | "value_extra"
ROWS: Dict[str, tuple] = {
    "distance_locking": (2, "value"),
    "hw_axis_falt": (8, "value_extra"),
    "dormas_cilindru": (9, "value"),
    "balama_ax": (11, "value_check"),
    "foarfeca": (12, "value_check"),
    "balama_simp_desc": (13, "value_check"),
    "balama_toc_jos": (14, "value_check"),
    "balama_ccv_jos": (15, "value_check"),
    "blocator_std": (17, "value"),
    "blocator_siguranta": (18, "value"),
    "blocator_oscilobatare": (19, "value"),
    "blocator_ob_rc": (20, "value"),
    "blocator_tbt": (21, "value"),
    "ridicator_cercevea": (22, "value"),
    "blocator_ob_bs": (23, "value"),
    "componente_ccv_bs": (24, "value"),
    "tip_cremon": (25, "value"),
    "mecanism_stulp": (26, "value"),
    "dfe": (27, "check"),
    "tfe": (28, "check"),
    "msl": (30, "check"),
    "zss": (31, "check"),
    "drive_rod_gak2450": (32, "check"),
    "support_plate": (33, "check"),
    "blocking_plate": (34, "check"),
    "pull_in_device": (35, "check"),
    "usi_dormas": (36, "value"),
    "broasca_chei": (37, "check"),
    "broasca_maner": (38, "check"),
    "broasca_siguranta": (39, "check"),
    "broasca_av4": (40, "check"),
    "blocator_wsk": (41, "value"),
    "usi_blocator_siguranta": (42, "value"),
    "contraplaca_centrala": (43, "value"),
    "contraplaca_sus_jos": (44, "value"),
    "contraplaca_continua": (45, "value"),
}


# ---------------------------------------------------------------- export
def _cell_re(ref: str) -> re.Pattern:
    return re.compile(r'<c r="' + ref + r'"(?P<attrs>[^>/]*)(?:/>|>.*?</c>)', re.S)


def _cell_xml(ref: str, style: str, value) -> str:
    """Serialise one <c> element for a bool (native checkbox), empty, or inline-string value."""
    if isinstance(value, bool):
        return f'<c r="{ref}"{style} t="b"><v>{1 if value else 0}</v></c>'
    if value is None or str(value) == "":
        return f'<c r="{ref}"{style}/>'
    return f'<c r="{ref}"{style} t="inlineStr"><is><t xml:space="preserve">{escape(str(value))}</t></is></c>'


def _set_cell(xml: str, ref: str, value) -> str:
    """Replace one <c> element; keeps its style (s) so formatting/checkbox rendering is preserved."""
    m = _cell_re(ref).search(xml)
    if not m:
        return xml  # template always has these cells; skip defensively
    s = re.search(r'\ss="(\d+)"', m.group("attrs"))
    style = f' s="{s.group(1)}"' if s else ""
    return xml[: m.start()] + _cell_xml(ref, style, value) + xml[m.end():]


def _write_colors(xml: str, vc: str, cc: str, cell: SolutionCell) -> str:
    colors = [c.strip() for c in cell.value.split(",") if c.strip()]
    slots = [f"{vc}{r}" for r in COLOR_ROWS] + [f"{cc}{r}" for r in COLOR_ROWS]
    for slot_i, ref in enumerate(slots):
        xml = _set_cell(xml, ref, colors[slot_i] if slot_i < len(colors) else "")
    return xml


def _write_row(xml: str, vc: str, cc: str, row: int, mode: str, cell: SolutionCell) -> str:
    if mode != "check":
        xml = _set_cell(xml, f"{vc}{row}", cell.value)
    if mode == "value_extra":
        xml = _set_cell(xml, f"{cc}{row}", cell.extra)
    if mode in ("value_check", "check"):
        xml = _set_cell(xml, f"{cc}{row}", cell.checked)
    return xml


def _write_system(xml: str, i: int, system: SolutionSystem) -> str:
    vc, cc = VALUE_COLS[i], CHECK_COLS[i]
    xml = _set_cell(xml, f"{vc}{NAME_ROW}", system.name)
    xml = _write_colors(xml, vc, cc, system.cells.get("cap_color", SolutionCell()))
    for key, (row, mode) in ROWS.items():
        xml = _write_row(xml, vc, cc, row, mode, system.cells.get(key, SolutionCell()))
    return xml


def _rezip(zin: zipfile.ZipFile, sheet_xml: str) -> bytes:
    """Copy every part of the template unchanged except the patched worksheet."""
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as zout:
        for item in zin.infolist():
            data = sheet_xml.encode("utf-8") if item.filename == SHEET_PATH else zin.read(item.filename)
            zout.writestr(item, data, compress_type=zipfile.ZIP_DEFLATED)
    return out.getvalue()


def export_solution(table: SolutionTable) -> bytes:
    with zipfile.ZipFile(TEMPLATE) as zin:
        xml = zin.read(SHEET_PATH).decode("utf-8")
        for i, system in enumerate(table.systems[: len(VALUE_COLS)]):
            xml = _write_system(xml, i, system)
        return _rezip(zin, xml)


# ---------------------------------------------------------------- import
def _text(v) -> str:
    if v is None or isinstance(v, bool):
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _truthy(v) -> bool:
    if isinstance(v, bool):
        return v
    return _text(v).lower() in ("true", "1", "x", "da", "yes", "✓")


def _load_sheet(data: bytes):
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    except Exception as e:  # noqa: BLE001 - surface any parse failure as a user error
        raise ValueError("Could not read the file — upload a .xlsx Solution sheet") from e
    ws = wb.active
    if not _text(ws["A1"].value).lower().startswith("profile system"):
        raise ValueError("This doesn't look like a Solution-EDV sheet (A1 should be 'Profile system:')")
    return ws


def _read_colors(ws, vc: str, cc: str) -> Optional[SolutionCell]:
    colors = [_text(ws[f"{col}{r}"].value) for col in (vc, cc) for r in COLOR_ROWS]
    colors = [c for c in colors if c]
    return SolutionCell(value=", ".join(colors)) if colors else None


def _read_row(ws, vc: str, cc: str, row: int, mode: str) -> Optional[SolutionCell]:
    value = _text(ws[f"{vc}{row}"].value) if mode != "check" else ""
    extra = _text(ws[f"{cc}{row}"].value) if mode == "value_extra" else ""
    checked = _truthy(ws[f"{cc}{row}"].value) if mode in ("value_check", "check") else False
    return SolutionCell(value=value, extra=extra, checked=checked) if (value or extra or checked) else None


def _is_template_default(key: str, cell: SolutionCell) -> bool:
    # the blank template pre-fills "25/30" in every column — not enough to count as a system
    return key == "dormas_cilindru" and cell.value == "25/30"


def _read_system(ws, i: int) -> Optional[SolutionSystem]:
    """One profile-system column, or None when the column is blank."""
    vc, cc = VALUE_COLS[i], CHECK_COLS[i]
    name = _text(ws[f"{vc}{NAME_ROW}"].value)
    cells: Dict[str, SolutionCell] = {}
    colors = _read_colors(ws, vc, cc)
    if colors:
        cells["cap_color"] = colors
    for key, (row, mode) in ROWS.items():
        cell = _read_row(ws, vc, cc, row, mode)
        if cell:
            cells[key] = cell
    meaningful = any(not _is_template_default(k, c) for k, c in cells.items())
    if not (name or meaningful):
        return None
    return SolutionSystem(name=name or f"System {i + 1}", cells=cells)


def import_solution(data: bytes) -> SolutionTable:
    """Parse a filled Solution-EDV workbook. Raises ValueError if unreadable or empty."""
    ws = _load_sheet(data)
    systems = [s for s in (_read_system(ws, i) for i in range(len(VALUE_COLS))) if s]
    if not systems:
        raise ValueError("No profile systems found in the sheet")
    return SolutionTable(systems=systems)


def safe_filename(company: str) -> str:
    base = re.sub(r"[^\w\- ]+", "", company).strip().replace(" ", "-") or "client"
    return f"Solution-{base}.xlsx"


def max_upload_bytes() -> int:
    return 5 * 1024 * 1024
